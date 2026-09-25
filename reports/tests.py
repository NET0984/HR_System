from datetime import date, datetime, time, timedelta
from io import BytesIO
import re

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from pypdf import PdfReader

from attendance.models import Attendance
from employees.models import Department, Employee
from leave.models import LeaveRequest


User = get_user_model()


def _aware(day, hour):
    return timezone.make_aware(datetime.combine(day, time(hour, 0)), timezone.get_current_timezone())


class MonthlyReportTests(TestCase):
    def setUp(self):
        self.hr = User.objects.create_user(
            username='hr-user',
            password='test-password-123',
            role=User.ROLE_HR_ADMIN,
        )
        self.employee_user = User.objects.create_user(
            username='emp-user',
            password='test-password-123',
            role=User.ROLE_EMPLOYEE,
        )
        self.department = Department.objects.create(name='Kỹ thuật')
        self.employee = Employee.objects.create(
            employee_code='NV0001',
            full_name='Nguyễn Minh Anh',
            joining_date=date(2024, 1, 15),
            department=self.department,
            user=self.employee_user,
        )
        today = timezone.localdate()
        self.month_start = today.replace(day=1)
        self.month = f'{today.year:04d}-{today.month:02d}'

    def _completed(self, day, in_hour, out_hour, revision=1, source=Attendance.SOURCE_BROWSER):
        return Attendance.objects.create(
            employee=self.employee,
            date=day,
            original_check_in=_aware(day, in_hour),
            original_check_out=_aware(day, out_hour),
            effective_check_in=_aware(day, in_hour),
            effective_check_out=_aware(day, out_hour),
            revision=revision,
            source=source,
        )

    def test_employee_cannot_view_report(self):
        self.client.login(username='emp-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'))

        self.assertRedirects(response, '/login/?next=/reports/')

    def test_hr_can_view_report(self):
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'), {'month': self.month})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Báo cáo chấm công tháng')

    def test_totals_count_completed_incomplete_and_corrected(self):
        self._completed(self.month_start, 8, 16)
        self._completed(self.month_start + timedelta(days=1), 8, 17, revision=2, source=Attendance.SOURCE_CORRECTION)
        Attendance.objects.create(
            employee=self.employee,
            date=self.month_start + timedelta(days=2),
            original_check_in=_aware(self.month_start + timedelta(days=2), 8),
            effective_check_in=_aware(self.month_start + timedelta(days=2), 8),
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'), {'month': self.month})

        totals = response.context['totals']
        self.assertEqual(totals['completed_days'], 2)
        self.assertEqual(totals['total_hours'], 17.0)
        self.assertEqual(totals['incomplete_days'], 1)
        self.assertEqual(totals['corrected_days'], 1)

    def test_approved_leave_days_are_counted_and_overlap_flagged(self):
        self._completed(self.month_start, 8, 16)
        LeaveRequest.objects.create(
            employee=self.employee,
            leave_type=LeaveRequest.TYPE_ANNUAL,
            start_date=self.month_start,
            end_date=self.month_start + timedelta(days=1),
            reason='Nghỉ phép',
            state=LeaveRequest.STATE_APPROVED,
            requester=self.employee_user,
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'), {'month': self.month})

        totals = response.context['totals']
        self.assertEqual(totals['leave_days'], 2)
        self.assertTrue(totals['overlap'])

    def test_employee_filter_returns_single_detail(self):
        self._completed(self.month_start, 8, 16)
        other = Employee.objects.create(
            employee_code='NV0002',
            full_name='Trần Quốc Bảo',
            joining_date=date(2024, 1, 15),
            department=self.department,
        )
        Attendance.objects.create(
            employee=other,
            date=self.month_start,
            original_check_in=_aware(self.month_start, 9),
            effective_check_in=_aware(self.month_start, 9),
            effective_check_out=_aware(self.month_start, 18),
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'), {
            'month': self.month,
            'employee': self.employee.pk,
        })

        rows = response.context['rows']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['employee'], self.employee)
        self.assertIsNotNone(response.context['single'])

    def test_department_filter(self):
        other_department = Department.objects.create(name='Kinh doanh')
        Employee.objects.create(
            employee_code='NV0002',
            full_name='Trần Quốc Bảo',
            joining_date=date(2024, 1, 15),
            department=other_department,
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:monthly'), {
            'month': self.month,
            'department': other_department.pk,
        })

        codes = [row['employee'].employee_code for row in response.context['rows']]
        self.assertEqual(codes, ['NV0002'])


class DashboardMetricsTests(TestCase):
    def setUp(self):
        self.hr = User.objects.create_user(
            username='hr-user',
            password='test-password-123',
            role=User.ROLE_HR_ADMIN,
        )
        self.employee_user = User.objects.create_user(
            username='emp-user',
            password='test-password-123',
            role=User.ROLE_EMPLOYEE,
        )
        self.unlinked_user = User.objects.create_user(
            username='no-profile',
            password='test-password-123',
            role=User.ROLE_EMPLOYEE,
        )
        self.employee = Employee.objects.create(
            employee_code='NV0001',
            full_name='Nguyễn Minh Anh',
            joining_date=date(2024, 1, 15),
            is_active=True,
            user=self.employee_user,
        )
        self.inactive_employee = Employee.objects.create(
            employee_code='NV0002',
            full_name='Trần Quốc Bảo',
            joining_date=date(2024, 1, 15),
            is_active=False,
        )

    def test_hr_dashboard_shows_company_metrics(self):
        today = timezone.localdate()
        Attendance.objects.create(
            employee=self.employee,
            date=today,
            original_check_in=timezone.now(),
            effective_check_in=timezone.now(),
        )
        from attendance.models import CorrectionRequest
        CorrectionRequest.objects.create(
            employee=self.employee,
            date=today,
            base_revision=1,
            proposed_check_in=timezone.now(),
            proposed_check_out=timezone.now(),
            reason='Sửa',
            state=CorrectionRequest.STATE_PENDING,
        )
        LeaveRequest.objects.create(
            employee=self.employee,
            leave_type=LeaveRequest.TYPE_ANNUAL,
            start_date=today,
            end_date=today,
            reason='Nghỉ',
            state=LeaveRequest.STATE_PENDING,
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('accounts:home'))

        metrics = response.context['metrics_company']
        self.assertEqual(metrics['employees'], 1)
        self.assertEqual(metrics['present_today'], 1)
        self.assertEqual(metrics['pending_corrections'], 1)
        self.assertEqual(metrics['pending_leave'], 1)
        self.assertNotIn('metrics_personal', response.context)
        self.assertContains(response, 'Có mặt hôm nay')
        self.assertNotContains(response, 'Ngày công tháng này')

    def test_employee_dashboard_shows_personal_metrics(self):
        from attendance.models import CorrectionRequest
        today = timezone.localdate()
        month_start = today.replace(day=1)
        Attendance.objects.create(
            employee=self.employee,
            date=month_start,
            original_check_in=_aware(month_start, 8),
            original_check_out=_aware(month_start, 17),
            effective_check_in=_aware(month_start, 8),
            effective_check_out=_aware(month_start, 17),
        )
        CorrectionRequest.objects.create(
            employee=self.employee,
            date=month_start,
            base_revision=1,
            proposed_check_in=timezone.now(),
            proposed_check_out=timezone.now(),
            reason='Sửa',
            state=CorrectionRequest.STATE_PENDING,
        )
        LeaveRequest.objects.create(
            employee=self.employee,
            leave_type=LeaveRequest.TYPE_ANNUAL,
            start_date=today,
            end_date=today,
            reason='Nghỉ',
            state=LeaveRequest.STATE_PENDING,
        )
        self.client.login(username='emp-user', password='test-password-123')

        response = self.client.get(reverse('accounts:home'))

        metrics = response.context['metrics_personal']
        self.assertEqual(metrics['completed_days'], 1)
        self.assertEqual(metrics['total_hours'], 9.0)
        self.assertEqual(metrics['pending_corrections'], 1)
        self.assertEqual(metrics['pending_leave'], 1)
        self.assertNotIn('metrics_company', response.context)
        self.assertContains(response, 'Ngày công tháng này')
        self.assertNotContains(response, 'Có mặt hôm nay')

    def test_employee_without_profile_hides_metric_row(self):
        self.client.login(username='no-profile', password='test-password-123')

        response = self.client.get(reverse('accounts:home'))

        self.assertNotIn('metrics_company', response.context)
        self.assertNotIn('metrics_personal', response.context)
        self.assertNotContains(response, 'Ngày công tháng này')
        self.assertNotContains(response, 'Có mặt hôm nay')


class ReportExportTests(TestCase):
    def setUp(self):
        self.hr = User.objects.create_user(
            username='hr-user',
            password='test-password-123',
            role=User.ROLE_HR_ADMIN,
        )
        self.employee_user = User.objects.create_user(
            username='emp-user',
            password='test-password-123',
            role=User.ROLE_EMPLOYEE,
        )
        self.department = Department.objects.create(name='Kỹ thuật')
        self.employee = Employee.objects.create(
            employee_code='NV0001',
            full_name='Nguyễn Minh Anh',
            joining_date=date(2024, 1, 15),
            department=self.department,
            user=self.employee_user,
        )
        self.today = timezone.localdate()
        self.month_start = self.today.replace(day=1)
        self.month = f'{self.today.year:04d}-{self.today.month:02d}'
        Attendance.objects.create(
            employee=self.employee,
            date=self.month_start,
            original_check_in=_aware(self.month_start, 8),
            original_check_out=_aware(self.month_start, 17),
            effective_check_in=_aware(self.month_start, 8),
            effective_check_out=_aware(self.month_start, 17),
        )

    def test_pdf_export_downloads_valid_pdf(self):
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:export_pdf'), {'month': self.month})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('.pdf', response['Content-Disposition'])
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertGreater(len(response.content), 2000)

    def test_pdf_export_contains_vietnamese_content(self):
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:export_pdf'), {'month': self.month})

        reader = PdfReader(BytesIO(response.content))
        text = re.sub(r'\s+', ' ', ''.join(page.extract_text() or '' for page in reader.pages))
        self.assertIn('BÁO CÁO CHẤM CÔNG THÁNG', text)
        self.assertIn('Nguyễn Minh Anh', text)
        self.assertIn('Kỹ thuật', text)
        self.assertIn('Tổng cộng', text)

    def test_excel_export_downloads_valid_workbook(self):
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:export_excel'), {'month': self.month})

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml.sheet', response['Content-Type'])
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('.xlsx', response['Content-Disposition'])

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        sheet = workbook.active
        values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
        self.assertIn('BÁO CÁO CHẤM CÔNG THÁNG', values)
        self.assertIn('Mã NV', values)
        self.assertIn('NV0001', values)
        self.assertIn('Nguyễn Minh Anh', values)
        self.assertIn('Tổng cộng', values)

    def test_excel_export_respects_employee_filter(self):
        other = Employee.objects.create(
            employee_code='NV0002',
            full_name='Trần Quốc Bảo',
            joining_date=date(2024, 1, 15),
            department=self.department,
        )
        self.client.login(username='hr-user', password='test-password-123')

        response = self.client.get(reverse('reports:export_excel'), {
            'month': self.month,
            'employee': other.pk,
        })

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        values = [str(cell.value) for row in workbook.active.iter_rows() for cell in row if cell.value is not None]
        self.assertIn('NV0002', values)
        self.assertNotIn('NV0001', values)

    def test_employee_cannot_export(self):
        self.client.login(username='emp-user', password='test-password-123')

        pdf_response = self.client.get(reverse('reports:export_pdf'))
        excel_response = self.client.get(reverse('reports:export_excel'))

        self.assertRedirects(pdf_response, '/login/?next=/reports/export/pdf/')
        self.assertRedirects(excel_response, '/login/?next=/reports/export/excel/')
