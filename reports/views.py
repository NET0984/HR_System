import os
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import openpyxl
from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import render
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

from accounts.views import hr_required
from attendance.models import Attendance
from employees.models import Department, Employee
from leave.models import LeaveRequest

from .pdf_support import enable_local_font_files


BUNDLED_FONT = Path(settings.BASE_DIR) / 'static' / 'fonts' / 'DejaVuSans.ttf'
FONT_CANDIDATES = [
    str(BUNDLED_FONT),
    r'C:\Windows\Fonts\arial.ttf',
    '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
]

EXCEL_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _month_options(today):
    options = []
    year, month = today.year, today.month
    for _ in range(12):
        options.append((f'{year:04d}-{month:02d}', f'Tháng {month:02d}/{year}'))
        month -= 1
        if month == 0:
            month = 12
            year -= 1
    return options


def _month_bounds(year, month):
    start = date(year, month, 1)
    next_month = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return start, next_month - timedelta(days=1)


def _pdf_font_path():
    for path in FONT_CANDIDATES:
        if os.path.exists(path):
            return path.replace('\\', '/')
    return None


def _pdf_font_bold_path():
    regular = _pdf_font_path()
    if not regular:
        return None
    candidate = Path(regular).with_name('DejaVuSans-Bold.ttf')
    if candidate.exists():
        return str(candidate).replace('\\', '/')
    return None


def _build_report_data(request):
    today = timezone.localdate()
    selected_month = request.GET.get('month', f'{today.year:04d}-{today.month:02d}')
    selected_department = request.GET.get('department', '')
    selected_employee = request.GET.get('employee', '')

    try:
        year, month = (int(part) for part in selected_month.split('-'))
        month_start, month_end = _month_bounds(year, month)
    except (ValueError, TypeError):
        year, month = today.year, today.month
        selected_month = f'{year:04d}-{month:02d}'
        month_start, month_end = _month_bounds(year, month)

    employees = Employee.objects.select_related('department').order_by('employee_code')
    if selected_department:
        employees = employees.filter(department_id=selected_department)
    if selected_employee:
        employees = employees.filter(pk=selected_employee)

    employee_ids = list(employees.values_list('pk', flat=True))

    attendance_records = Attendance.objects.filter(
        employee_id__in=employee_ids,
        date__gte=month_start,
        date__lte=month_end,
    )
    leave_requests = LeaveRequest.objects.filter(
        employee_id__in=employee_ids,
        state=LeaveRequest.STATE_APPROVED,
        start_date__lte=month_end,
        end_date__gte=month_start,
    )

    attendance_by_employee = {}
    for record in attendance_records:
        attendance_by_employee.setdefault(record.employee_id, []).append(record)

    leave_by_employee = {}
    for leave in leave_requests:
        leave_by_employee.setdefault(leave.employee_id, []).append(leave)

    rows = []
    for employee in employees:
        records = attendance_by_employee.get(employee.pk, [])
        completed = [r for r in records if r.effective_check_in and r.effective_check_out]
        incomplete = [r for r in records if r.effective_check_in and not r.effective_check_out]
        corrected = [r for r in records if r.revision > 1 or r.source == Attendance.SOURCE_CORRECTION]
        total_seconds = sum(
            (r.effective_check_out - r.effective_check_in).total_seconds() for r in completed
        )

        record_dates = {r.date for r in records}
        leave_days = 0
        overlap = False
        for leave in leave_by_employee.get(employee.pk, []):
            start = max(leave.start_date, month_start)
            end = min(leave.end_date, month_end)
            if end >= start:
                leave_days += (end - start).days + 1
                if any(start <= day <= end for day in record_dates):
                    overlap = True

        rows.append({
            'employee': employee,
            'completed_days': len(completed),
            'total_hours': round(total_seconds / 3600, 1),
            'incomplete_days': len(incomplete),
            'corrected_days': len(corrected),
            'leave_days': leave_days,
            'overlap': overlap,
            'records': sorted(records, key=lambda item: item.date),
        })

    totals = {
        'completed_days': sum(row['completed_days'] for row in rows),
        'total_hours': round(sum(row['total_hours'] for row in rows), 1),
        'incomplete_days': sum(row['incomplete_days'] for row in rows),
        'corrected_days': sum(row['corrected_days'] for row in rows),
        'leave_days': sum(row['leave_days'] for row in rows),
        'overlap': any(row['overlap'] for row in rows),
    }

    leave_date_strings = []
    single = rows[0] if selected_employee and len(rows) == 1 else None
    if single:
        for leave in leave_by_employee.get(single['employee'].pk, []):
            start = max(leave.start_date, month_start)
            end = min(leave.end_date, month_end)
            day = start
            while day <= end:
                leave_date_strings.append(day.isoformat())
                day += timedelta(days=1)

    return {
        'selected_month': selected_month,
        'selected_department': selected_department,
        'selected_employee': selected_employee,
        'month_options': _month_options(today),
        'departments': Department.objects.all(),
        'employees': Employee.objects.order_by('employee_code'),
        'rows': rows,
        'totals': totals,
        'chart_labels': [row['employee'].full_name for row in rows],
        'chart_values': [row['total_hours'] for row in rows],
        'single': single,
        'leave_date_strings': leave_date_strings,
        'month_start': month_start,
        'month_end': month_end,
        'font_path': _pdf_font_path(),
        'font_path_bold': _pdf_font_bold_path(),
    }


@hr_required
def monthly_report(request):
    return render(request, 'reports/monthly.html', _build_report_data(request))


@hr_required
def monthly_report_pdf(request):
    enable_local_font_files()
    data = _build_report_data(request)
    html = render_to_string('reports/monthly_pdf.html', data)
    buffer = BytesIO()
    result = pisa.CreatePDF(html, dest=buffer, encoding='utf-8')
    if result.err:
        return HttpResponse('Không tạo được PDF.', status=500)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    filename = f'bao-cao-cham-cong-{data["selected_month"]}.pdf'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@hr_required
def monthly_report_excel(request):
    data = _build_report_data(request)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Bao cao cham cong'

    bold = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='DDEBF7')
    center = Alignment(horizontal='center')

    sheet['A1'] = 'BÁO CÁO CHẤM CÔNG THÁNG'
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A2'] = f'Kỳ báo cáo: {data["month_start"]:%d/%m/%Y} - {data["month_end"]:%d/%m/%Y}'
    sheet['A3'] = f'Tháng: {data["selected_month"]}'

    headers = ['Mã NV', 'Nhân viên', 'Phòng ban', 'Ngày công', 'Tổng giờ',
               'Thiếu giờ ra', 'Ngày được sửa', 'Nghỉ phép', 'Ghi chú']
    header_row = 5
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center

    current_row = header_row + 1
    for row in data['rows']:
        sheet.cell(row=current_row, column=1, value=row['employee'].employee_code)
        sheet.cell(row=current_row, column=2, value=row['employee'].full_name)
        sheet.cell(row=current_row, column=3, value=row['employee'].department.name if row['employee'].department else '')
        sheet.cell(row=current_row, column=4, value=row['completed_days'])
        sheet.cell(row=current_row, column=5, value=row['total_hours'])
        sheet.cell(row=current_row, column=6, value=row['incomplete_days'])
        sheet.cell(row=current_row, column=7, value=row['corrected_days'])
        sheet.cell(row=current_row, column=8, value=row['leave_days'])
        sheet.cell(row=current_row, column=9, value='Trùng nghỉ phép' if row['overlap'] else '')
        current_row += 1

    totals = data['totals']
    sheet.cell(row=current_row, column=1, value='Tổng cộng').font = bold
    for column, value in (
        (4, totals['completed_days']),
        (5, totals['total_hours']),
        (6, totals['incomplete_days']),
        (7, totals['corrected_days']),
        (8, totals['leave_days']),
    ):
        cell = sheet.cell(row=current_row, column=column, value=value)
        cell.font = bold

    widths = [12, 26, 18, 12, 12, 14, 16, 12, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)

    response = HttpResponse(buffer.getvalue(), content_type=EXCEL_CONTENT_TYPE)
    filename = f'bao-cao-cham-cong-{data["selected_month"]}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
